#!/usr/bin/env python3
"""
Jira Extraction Application
============================
Extracts issue data from Jira using the REST API and generates
formatted Excel reports for analysis and stakeholder communication.

Usage:
    python jira_extraction.py                         # Default: all issues
    python jira_extraction.py --jql "project = PROJ"  # Custom JQL query
    python jira_extraction.py --template open_bugs    # Use a saved template
    python jira_extraction.py --limit 50              # Limit results
"""

import argparse
import json
import logging
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from jira import JIRA, JIRAError
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
LOG_FORMAT = "%(asctime)s [%(levelname)s] %(message)s"
logging.basicConfig(level=logging.INFO, format=LOG_FORMAT)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Configuration helpers
# ---------------------------------------------------------------------------

def load_config(config_path: str = "config.json") -> dict:
    """Load application configuration from a JSON file."""
    path = Path(config_path)
    if not path.exists():
        logger.warning("Config file '%s' not found – using defaults.", config_path)
        return {
            "default_jql": "project IS NOT EMPTY ORDER BY created DESC",
            "max_results_per_page": 100,
            "output_directory": "output",
            "fields": [],
            "report_templates": {},
        }
    with open(path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def load_credentials() -> tuple:
    """Load Jira credentials from environment / .env file."""
    load_dotenv()
    url = os.getenv("JIRA_URL")
    email = os.getenv("JIRA_EMAIL")
    token = os.getenv("JIRA_API_TOKEN")

    missing = []
    if not url:
        missing.append("JIRA_URL")
    if not email:
        missing.append("JIRA_EMAIL")
    if not token:
        missing.append("JIRA_API_TOKEN")

    if missing:
        logger.error(
            "Missing required environment variables: %s. "
            "Please set them in a .env file or your shell environment.",
            ", ".join(missing),
        )
        sys.exit(1)

    return url, email, token


# ---------------------------------------------------------------------------
# Jira interaction
# ---------------------------------------------------------------------------

def connect_to_jira(url: str, email: str, token: str) -> JIRA:
    """Authenticate and return a JIRA client instance."""
    try:
        jira = JIRA(server=url, basic_auth=(email, token))
        server_info = jira.server_info()
        logger.info(
            "Connected to Jira  –  Server: %s  |  Version: %s",
            server_info.get("baseUrl", url),
            server_info.get("version", "unknown"),
        )
        return jira
    except JIRAError as exc:
        logger.error("Failed to connect to Jira: %s", exc.text if hasattr(exc, 'text') else exc)
        sys.exit(1)
    except Exception as exc:
        logger.error("Unexpected error connecting to Jira: %s", exc)
        sys.exit(1)


class _IssueProxy:
    """Lightweight wrapper so JSON dicts look like jira.Issue objects."""

    def __init__(self, data: dict):
        self.key = data["key"]
        self.raw = data
        self.fields = _FieldsProxy(data.get("fields", {}))


class _FieldsProxy:
    """Exposes Jira field dicts as attributes (with nested name access)."""

    def __init__(self, fields: dict):
        self._fields = fields
        for k, v in fields.items():
            setattr(self, k, self._wrap(v))

    @staticmethod
    def _wrap(value):
        if isinstance(value, dict):
            return _NameAccessor(value)
        return value


class _NameAccessor:
    """Allow obj.name / obj.displayName / str(obj) on plain dicts."""

    def __init__(self, d: dict):
        self._d = d

    def __getattr__(self, name):
        try:
            return self._d[name]
        except KeyError:
            return None

    def __str__(self):
        return self._d.get("name") or self._d.get("displayName") or str(self._d)


def fetch_issues(url: str, email: str, token: str, jql: str,
                 max_per_page: int = 100,
                 limit: int | None = None) -> list:
    """Retrieve all issues matching *jql* via the v3 search/jql endpoint."""
    import requests as _req
    from requests.auth import HTTPBasicAuth

    all_issues: list[_IssueProxy] = []
    start_at = 0
    total = None
    endpoint = f"{url.rstrip('/')}/rest/api/3/search/jql"
    auth = HTTPBasicAuth(email, token)

    logger.info("Executing JQL: %s", jql)

    while True:
        batch_size = max_per_page
        if limit is not None:
            remaining = limit - len(all_issues)
            if remaining <= 0:
                break
            batch_size = min(batch_size, remaining)

        params = {
            "jql": jql,
            "startAt": start_at,
            "maxResults": batch_size,
            "fields": "*all",
        }

        try:
            resp = _req.get(endpoint, params=params, auth=auth, timeout=30)
            resp.raise_for_status()
            data = resp.json()
        except _req.exceptions.HTTPError as exc:
            body = exc.response.text if exc.response is not None else str(exc)
            logger.error("JQL query failed (%s): %s", exc.response.status_code if exc.response else "?", body)
            sys.exit(1)
        except Exception as exc:
            logger.error("JQL query failed: %s", exc)
            sys.exit(1)

        issues_batch = [_IssueProxy(i) for i in data.get("issues", [])]

        if total is None:
            total = data.get("total", len(issues_batch))
            effective_total = min(total, limit) if limit else total
            logger.info("Total issues matching query: %d (fetching %d)", total, effective_total)

        all_issues.extend(issues_batch)
        logger.info("  Fetched %d / %d issues …", len(all_issues),
                     min(total, limit) if limit else total)

        if len(issues_batch) < batch_size or len(all_issues) >= total:
            break
        if limit and len(all_issues) >= limit:
            break

        start_at += len(issues_batch)

    logger.info("Retrieved %d issues in total.", len(all_issues))
    return all_issues


# ---------------------------------------------------------------------------
# Data transformation
# ---------------------------------------------------------------------------

def _safe(value, default=""):
    """Return *value* or *default* when None."""
    return value if value is not None else default


def _fmt_datetime(dt_str: str | None) -> str:
    """Convert Jira datetime string to a human-readable format."""
    if not dt_str:
        return ""
    try:
        dt = datetime.fromisoformat(dt_str.replace("Z", "+00:00"))
        return dt.strftime("%Y-%m-%d %H:%M")
    except (ValueError, TypeError):
        return str(dt_str)


def _seconds_to_hours(seconds) -> str:
    """Convert seconds (int/None) to a readable hours string."""
    if seconds is None:
        return ""
    try:
        hours = int(seconds) / 3600
        return f"{hours:.1f}h"
    except (ValueError, TypeError):
        return ""


def _get_sprint_name(fields) -> str:
    """Extract the sprint name from the issue fields."""
    sprint = getattr(fields, "sprint", None)
    if sprint is None:
        # Try customfield for sprint
        for attr in dir(fields):
            if attr.startswith("customfield_"):
                val = getattr(fields, attr, None)
                if isinstance(val, list):
                    for item in val:
                        if isinstance(item, str) and "name=" in item:
                            # Parse sprint string format
                            try:
                                name_part = item.split("name=")[1].split(",")[0]
                                return name_part
                            except (IndexError, AttributeError):
                                pass
                        elif hasattr(item, "name"):
                            return item.name
        return ""
    if hasattr(sprint, "name"):
        return sprint.name
    return str(sprint)


def issues_to_dataframe(issues: list) -> pd.DataFrame:
    """Convert a list of Jira issue objects into a pandas DataFrame."""
    records = []
    for issue in issues:
        f = issue.fields
        created_str = _fmt_datetime(getattr(f, "created", None))
        updated_str = _fmt_datetime(getattr(f, "updated", None))
        resolved_str = _fmt_datetime(getattr(f, "resolutiondate", None))

        # Calculate age in days
        age_days = ""
        if getattr(f, "created", None):
            try:
                created_dt = datetime.fromisoformat(
                    f.created.replace("Z", "+00:00")
                )
                age_days = (datetime.now(timezone.utc) - created_dt).days
            except (ValueError, TypeError):
                pass

        records.append(
            {
                "Key": issue.key,
                "Summary": _safe(getattr(f, "summary", None)),
                "Status": str(f.status) if getattr(f, "status", None) else "",
                "Priority": str(f.priority) if getattr(f, "priority", None) else "",
                "Issue Type": str(f.issuetype) if getattr(f, "issuetype", None) else "",
                "Assignee": (
                    f.assignee.displayName
                    if getattr(f, "assignee", None)
                    else "Unassigned"
                ),
                "Reporter": (
                    f.reporter.displayName
                    if getattr(f, "reporter", None)
                    else ""
                ),
                "Created": created_str,
                "Updated": updated_str,
                "Resolved": resolved_str,
                "Resolution": (
                    str(f.resolution)
                    if getattr(f, "resolution", None)
                    else ""
                ),
                "Labels": ", ".join(f.labels) if getattr(f, "labels", None) else "",
                "Components": (
                    ", ".join(c.name for c in f.components)
                    if getattr(f, "components", None)
                    else ""
                ),
                "Fix Versions": (
                    ", ".join(v.name for v in f.fixVersions)
                    if getattr(f, "fixVersions", None)
                    else ""
                ),
                "Sprint": _get_sprint_name(f),
                "Original Estimate": _seconds_to_hours(
                    getattr(f, "timeoriginalestimate", None)
                ),
                "Remaining Estimate": _seconds_to_hours(
                    getattr(f, "timeestimate", None)
                ),
                "Time Spent": _seconds_to_hours(
                    getattr(f, "timespent", None)
                ),
                "Age (Days)": age_days,
            }
        )

    df = pd.DataFrame(records)
    return df


# ---------------------------------------------------------------------------
# Excel generation
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="top")


def _auto_column_width(ws, df: pd.DataFrame, min_width: int = 10, max_width: int = 45):
    """Resize columns based on content."""
    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = len(str(col_name))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = (
            max(min_width, max_len + 3)
        )


def _style_header(ws, col_count: int):
    """Apply header styling to the first row."""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_data(ws, row_count: int, col_count: int):
    """Apply styling to data cells."""
    alt_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    for row in range(2, row_count + 2):  # +2 because of header at row 1
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = WRAP_ALIGNMENT
            if row % 2 == 0:
                cell.fill = alt_fill


def _create_summary_sheet(wb, df: pd.DataFrame):
    """Create a Summary worksheet with key metrics."""
    ws = wb.create_sheet("Summary")

    title_font = Font(name="Calibri", bold=True, size=14, color="1F4E79")
    section_font = Font(name="Calibri", bold=True, size=12, color="1F4E79")
    label_font = Font(name="Calibri", bold=True, size=10)
    value_font = Font(name="Calibri", size=10)
    section_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

    # Title
    ws.merge_cells("A1:D1")
    ws["A1"] = "Jira Extraction Report – Summary"
    ws["A1"].font = title_font

    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = value_font

    row = 4

    # --- Overview ---
    ws.cell(row=row, column=1, value="Overview").font = section_font
    ws.cell(row=row, column=1).fill = section_fill
    ws.cell(row=row, column=2).fill = section_fill
    row += 1
    ws.cell(row=row, column=1, value="Total Issues").font = label_font
    ws.cell(row=row, column=2, value=len(df)).font = value_font
    row += 2

    # --- Issues by Status ---
    if "Status" in df.columns:
        ws.cell(row=row, column=1, value="Issues by Status").font = section_font
        ws.cell(row=row, column=1).fill = section_fill
        ws.cell(row=row, column=2).fill = section_fill
        row += 1
        status_counts = df["Status"].value_counts()
        for status, count in status_counts.items():
            ws.cell(row=row, column=1, value=str(status)).font = label_font
            ws.cell(row=row, column=2, value=count).font = value_font
            row += 1
        row += 1

    # --- Issues by Priority ---
    if "Priority" in df.columns:
        ws.cell(row=row, column=1, value="Issues by Priority").font = section_font
        ws.cell(row=row, column=1).fill = section_fill
        ws.cell(row=row, column=2).fill = section_fill
        row += 1
        priority_counts = df["Priority"].value_counts()
        for priority, count in priority_counts.items():
            ws.cell(row=row, column=1, value=str(priority)).font = label_font
            ws.cell(row=row, column=2, value=count).font = value_font
            row += 1
        row += 1

    # --- Issues by Issue Type ---
    if "Issue Type" in df.columns:
        ws.cell(row=row, column=1, value="Issues by Type").font = section_font
        ws.cell(row=row, column=1).fill = section_fill
        ws.cell(row=row, column=2).fill = section_fill
        row += 1
        type_counts = df["Issue Type"].value_counts()
        for itype, count in type_counts.items():
            ws.cell(row=row, column=1, value=str(itype)).font = label_font
            ws.cell(row=row, column=2, value=count).font = value_font
            row += 1
        row += 1

    # --- Issues by Assignee (top 15) ---
    if "Assignee" in df.columns:
        ws.cell(row=row, column=1, value="Issues by Assignee (Top 15)").font = section_font
        ws.cell(row=row, column=1).fill = section_fill
        ws.cell(row=row, column=2).fill = section_fill
        row += 1
        assignee_counts = df["Assignee"].value_counts().head(15)
        for assignee, count in assignee_counts.items():
            ws.cell(row=row, column=1, value=str(assignee)).font = label_font
            ws.cell(row=row, column=2, value=count).font = value_font
            row += 1
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15


def generate_excel(df: pd.DataFrame, output_path: str) -> str:
    """Write the DataFrame to a formatted Excel workbook."""
    # Ensure output directory exists
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Issues")
        wb = writer.book
        ws = wb["Issues"]

        # Styling
        col_count = len(df.columns)
        row_count = len(df)

        _style_header(ws, col_count)
        _style_data(ws, row_count, col_count)
        _auto_column_width(ws, df)

        # Freeze the header row
        ws.freeze_panes = "A2"

        # Auto-filter
        ws.auto_filter.ref = ws.dimensions

        # Summary sheet
        _create_summary_sheet(wb, df)

    logger.info("Excel report saved to: %s", output_path)
    return output_path


# ---------------------------------------------------------------------------
# CLI & main
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(
        description="Jira Extraction Application – export Jira issues to Excel.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            '  python jira_extraction.py --jql "project = PROJ"\n'
            "  python jira_extraction.py --template open_bugs\n"
            "  python jira_extraction.py --limit 50\n"
            "  python jira_extraction.py --output my_report.xlsx\n"
        ),
    )
    parser.add_argument(
        "--jql",
        type=str,
        default=None,
        help="JQL query string. Overrides the default query from config.json.",
    )
    parser.add_argument(
        "--template",
        type=str,
        default=None,
        help="Name of a saved JQL template from config.json (e.g. open_bugs, recent_30_days).",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Maximum number of issues to retrieve.",
    )
    parser.add_argument(
        "--output",
        type=str,
        default=None,
        help="Output Excel file path. Defaults to output/jira_report_<timestamp>.xlsx.",
    )
    parser.add_argument(
        "--config",
        type=str,
        default="config.json",
        help="Path to config.json (default: config.json in current directory).",
    )
    return parser.parse_args()


def main():
    """Entry point."""
    args = parse_args()
    config = load_config(args.config)

    # --- Resolve JQL query ---
    if args.jql:
        jql = args.jql
    elif args.template:
        templates = config.get("report_templates", {})
        if args.template not in templates:
            logger.error(
                "Template '%s' not found. Available: %s",
                args.template,
                ", ".join(templates.keys()) or "(none)",
            )
            sys.exit(1)
        jql = templates[args.template]
        logger.info("Using template '%s'", args.template)
    else:
        jql = config.get("default_jql", "project IS NOT EMPTY ORDER BY created DESC")

    # --- Connect ---
    url, email, token = load_credentials()
    jira = connect_to_jira(url, email, token)

    # --- Fetch ---
    max_per_page = config.get("max_results_per_page", 100)
    issues = fetch_issues(url, email, token, jql, max_per_page=max_per_page, limit=args.limit)

    if not issues:
        logger.warning("No issues found for the given query. No report generated.")
        sys.exit(0)

    # --- Transform ---
    df = issues_to_dataframe(issues)
    logger.info("DataFrame shape: %d rows × %d columns", *df.shape)

    # --- Generate Excel ---
    if args.output:
        output_path = args.output
    else:
        out_dir = config.get("output_directory", "output")
        timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        output_path = os.path.join(out_dir, f"jira_report_{timestamp}.xlsx")

    generate_excel(df, output_path)

    print(f"\n{'='*60}")
    print(f"  [OK]  Jira Extraction Complete!")
    print(f"  Report: {os.path.abspath(output_path)}")
    print(f"  Issues: {len(df)}")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    main()
