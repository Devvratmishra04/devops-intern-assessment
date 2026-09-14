#!/usr/bin/env python3
"""
Jira Extraction Application – Automated Test Suite
====================================================
Comprehensive tests covering configuration loading, credential parsing,
data transformation, DataFrame schema mapping, Excel generation &
formatting, API pagination simulation, and live integration checks.

Run with:
    python run_tests.py              (from project root)
    python -m pytest tests/ -v       (if pytest installed)
"""

import json
import os
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

# Add parent directory to path so we can import jira_extraction
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import jira_extraction as je


# ---------------------------------------------------------------------------
# 1. Configuration & Environment Tests
# ---------------------------------------------------------------------------
class TestConfiguration(unittest.TestCase):
    """Tests for config.json loading and schema validation."""

    def test_load_config_default(self):
        """Config loads successfully from the project config.json."""
        config = je.load_config(
            str(Path(__file__).resolve().parent.parent / "config.json")
        )
        self.assertIn("default_jql", config)
        self.assertIn("max_results_per_page", config)
        self.assertIn("output_directory", config)
        self.assertIn("fields", config)
        self.assertIn("report_templates", config)

    def test_config_default_jql_is_string(self):
        config = je.load_config(
            str(Path(__file__).resolve().parent.parent / "config.json")
        )
        self.assertIsInstance(config["default_jql"], str)
        self.assertGreater(len(config["default_jql"]), 0)

    def test_config_max_results_range(self):
        config = je.load_config(
            str(Path(__file__).resolve().parent.parent / "config.json")
        )
        self.assertGreater(config["max_results_per_page"], 0)
        self.assertLessEqual(config["max_results_per_page"], 100)

    def test_config_fields_list(self):
        config = je.load_config(
            str(Path(__file__).resolve().parent.parent / "config.json")
        )
        self.assertIsInstance(config["fields"], list)
        self.assertGreater(len(config["fields"]), 0)
        # Must include essential fields
        for field in ["key", "summary", "status", "priority"]:
            self.assertIn(field, config["fields"])

    def test_config_templates_dict(self):
        config = je.load_config(
            str(Path(__file__).resolve().parent.parent / "config.json")
        )
        templates = config["report_templates"]
        self.assertIsInstance(templates, dict)
        expected_templates = ["all_issues", "open_bugs", "recent_30_days"]
        for name in expected_templates:
            self.assertIn(name, templates)
            self.assertIsInstance(templates[name], str)

    def test_load_config_missing_file_returns_defaults(self):
        config = je.load_config("nonexistent_config_file.json")
        self.assertIn("default_jql", config)
        self.assertEqual(config["max_results_per_page"], 100)

    def test_config_json_valid_syntax(self):
        """config.json is valid JSON."""
        config_path = Path(__file__).resolve().parent.parent / "config.json"
        with open(config_path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        self.assertIsInstance(data, dict)


class TestCredentials(unittest.TestCase):
    """Tests for .env credential parsing."""

    def test_env_file_exists(self):
        env_path = Path(__file__).resolve().parent.parent / ".env"
        self.assertTrue(env_path.exists(), ".env file must exist for deployment")

    def test_env_file_has_required_vars(self):
        env_path = Path(__file__).resolve().parent.parent / ".env"
        content = env_path.read_text()
        for var in ["JIRA_URL", "JIRA_EMAIL", "JIRA_API_TOKEN"]:
            self.assertIn(var, content, f"Missing {var} in .env")

    def test_env_file_no_placeholder_token(self):
        """Ensure the API token is not still a placeholder."""
        env_path = Path(__file__).resolve().parent.parent / ".env"
        content = env_path.read_text()
        self.assertNotIn("PASTE_YOUR_API_TOKEN_HERE", content)
        self.assertNotIn("your-real-api-token", content)


# ---------------------------------------------------------------------------
# 2. Data Transformation Tests
# ---------------------------------------------------------------------------
class TestDataTransformations(unittest.TestCase):
    """Tests for helper functions that parse and transform Jira data."""

    def test_safe_with_value(self):
        self.assertEqual(je._safe("hello"), "hello")
        self.assertEqual(je._safe(42), 42)

    def test_safe_with_none(self):
        self.assertEqual(je._safe(None), "")
        self.assertEqual(je._safe(None, "N/A"), "N/A")

    def test_fmt_datetime_valid_iso(self):
        result = je._fmt_datetime("2026-01-15T10:30:00.000+0000")
        self.assertIn("2026-01-15", result)
        self.assertIn("10:30", result)

    def test_fmt_datetime_with_z_suffix(self):
        result = je._fmt_datetime("2026-06-01T08:00:00Z")
        self.assertIn("2026-06-01", result)

    def test_fmt_datetime_none(self):
        self.assertEqual(je._fmt_datetime(None), "")

    def test_fmt_datetime_empty(self):
        self.assertEqual(je._fmt_datetime(""), "")

    def test_fmt_datetime_invalid(self):
        result = je._fmt_datetime("not-a-date")
        self.assertEqual(result, "not-a-date")

    def test_seconds_to_hours_normal(self):
        self.assertEqual(je._seconds_to_hours(3600), "1.0h")
        self.assertEqual(je._seconds_to_hours(7200), "2.0h")
        self.assertEqual(je._seconds_to_hours(5400), "1.5h")

    def test_seconds_to_hours_zero(self):
        self.assertEqual(je._seconds_to_hours(0), "0.0h")

    def test_seconds_to_hours_none(self):
        self.assertEqual(je._seconds_to_hours(None), "")

    def test_seconds_to_hours_invalid(self):
        self.assertEqual(je._seconds_to_hours("abc"), "")


class TestSprintExtraction(unittest.TestCase):
    """Tests for sprint name parsing from various field formats."""

    def test_sprint_from_name_attribute(self):
        """Sprint object with .name attribute."""
        fields = MagicMock()
        fields.sprint = MagicMock()
        fields.sprint.name = "Sprint 42"
        result = je._get_sprint_name(fields)
        self.assertEqual(result, "Sprint 42")

    def test_sprint_none(self):
        """No sprint set – should return empty string."""
        fields = MagicMock(spec=[])
        fields.sprint = None
        # No customfield_ attributes
        result = je._get_sprint_name(fields)
        self.assertEqual(result, "")

    def test_sprint_as_string(self):
        """Sprint returned as plain string (no .name)."""
        fields = MagicMock(spec=[])
        fields.sprint = "Sprint 99"
        result = je._get_sprint_name(fields)
        self.assertIn("Sprint 99", result)


class TestProxyObjects(unittest.TestCase):
    """Tests for _IssueProxy, _FieldsProxy, and _NameAccessor wrappers."""

    def test_issue_proxy_key(self):
        data = {"key": "PROJ-123", "fields": {"summary": "Test issue"}}
        proxy = je._IssueProxy(data)
        self.assertEqual(proxy.key, "PROJ-123")

    def test_issue_proxy_fields(self):
        data = {
            "key": "PROJ-1",
            "fields": {
                "summary": "My summary",
                "status": {"name": "Open"},
                "priority": {"name": "High"},
            },
        }
        proxy = je._IssueProxy(data)
        self.assertEqual(proxy.fields.summary, "My summary")
        self.assertEqual(str(proxy.fields.status), "Open")
        self.assertEqual(str(proxy.fields.priority), "High")

    def test_name_accessor_name(self):
        acc = je._NameAccessor({"name": "Done", "id": "3"})
        self.assertEqual(acc.name, "Done")
        self.assertEqual(str(acc), "Done")

    def test_name_accessor_display_name(self):
        acc = je._NameAccessor({"displayName": "John Doe"})
        self.assertEqual(acc.displayName, "John Doe")
        self.assertEqual(str(acc), "John Doe")

    def test_name_accessor_missing_key(self):
        acc = je._NameAccessor({"name": "Test"})
        self.assertIsNone(acc.nonexistent_key)


# ---------------------------------------------------------------------------
# 3. DataFrame Schema Tests
# ---------------------------------------------------------------------------
class TestDataFrameSchema(unittest.TestCase):
    """Tests for issues_to_dataframe column schema and edge cases."""

    def _make_issue(self, key="TEST-1", summary="Test", status="Open",
                    priority="Medium", issuetype="Task", assignee=None,
                    reporter=None, created=None, updated=None):
        """Create a mock _IssueProxy for testing."""
        fields = {
            "summary": summary,
            "status": {"name": status},
            "priority": {"name": priority},
            "issuetype": {"name": issuetype},
            "assignee": {"displayName": assignee} if assignee else None,
            "reporter": {"displayName": reporter} if reporter else None,
            "created": created or "2026-01-01T00:00:00.000+0000",
            "updated": updated or "2026-01-02T00:00:00.000+0000",
            "resolutiondate": None,
            "resolution": None,
            "labels": ["bug", "urgent"],
            "components": [],
            "fixVersions": [],
            "sprint": None,
            "timeoriginalestimate": 7200,
            "timeestimate": 3600,
            "timespent": 1800,
        }
        return je._IssueProxy({"key": key, "fields": fields})

    def test_dataframe_columns(self):
        """DataFrame has all expected columns."""
        issues = [self._make_issue()]
        df = je.issues_to_dataframe(issues)
        expected_cols = [
            "Key", "Summary", "Status", "Priority", "Issue Type",
            "Assignee", "Reporter", "Created", "Updated", "Resolved",
            "Resolution", "Labels", "Components", "Fix Versions",
            "Sprint", "Original Estimate", "Remaining Estimate",
            "Time Spent", "Age (Days)",
        ]
        for col in expected_cols:
            self.assertIn(col, df.columns, f"Missing column: {col}")

    def test_dataframe_row_count(self):
        issues = [self._make_issue(key=f"T-{i}") for i in range(5)]
        df = je.issues_to_dataframe(issues)
        self.assertEqual(len(df), 5)

    def test_unassigned_handling(self):
        """Unassigned issues should show 'Unassigned'."""
        issues = [self._make_issue(assignee=None)]
        df = je.issues_to_dataframe(issues)
        self.assertEqual(df.iloc[0]["Assignee"], "Unassigned")

    def test_label_formatting(self):
        """Labels should be comma-separated."""
        issues = [self._make_issue()]
        df = je.issues_to_dataframe(issues)
        self.assertIn("bug", df.iloc[0]["Labels"])
        self.assertIn("urgent", df.iloc[0]["Labels"])

    def test_time_estimates_formatted(self):
        """Time fields should be converted to hours."""
        issues = [self._make_issue()]
        df = je.issues_to_dataframe(issues)
        self.assertEqual(df.iloc[0]["Original Estimate"], "2.0h")
        self.assertEqual(df.iloc[0]["Remaining Estimate"], "1.0h")
        self.assertEqual(df.iloc[0]["Time Spent"], "0.5h")

    def test_age_days_calculated(self):
        """Age (Days) should be a positive integer."""
        issues = [self._make_issue(created="2026-01-01T00:00:00.000+0000")]
        df = je.issues_to_dataframe(issues)
        age = df.iloc[0]["Age (Days)"]
        self.assertTrue(
            isinstance(age, (int, float)) or hasattr(age, "__int__"),
            f"Age should be numeric, got {type(age)}"
        )
        self.assertGreater(int(age), 0)

    def test_empty_issues_list(self):
        df = je.issues_to_dataframe([])
        self.assertEqual(len(df), 0)


# ---------------------------------------------------------------------------
# 4. Excel Generation & Formatting Tests
# ---------------------------------------------------------------------------
class TestExcelGeneration(unittest.TestCase):
    """Tests for Excel workbook generation and styling."""

    def setUp(self):
        """Create a test DataFrame and temp output path."""
        self.test_data = [
            {
                "Key": "PROJ-1", "Summary": "First issue", "Status": "Open",
                "Priority": "High", "Issue Type": "Bug", "Assignee": "Alice",
                "Reporter": "Bob", "Created": "2026-01-01 00:00",
                "Updated": "2026-01-02 00:00", "Resolved": "",
                "Resolution": "", "Labels": "bug, critical",
                "Components": "", "Fix Versions": "", "Sprint": "Sprint 1",
                "Original Estimate": "4.0h", "Remaining Estimate": "2.0h",
                "Time Spent": "2.0h", "Age (Days)": 257,
            },
            {
                "Key": "PROJ-2", "Summary": "Second issue", "Status": "Done",
                "Priority": "Low", "Issue Type": "Story", "Assignee": "Charlie",
                "Reporter": "Alice", "Created": "2026-02-01 00:00",
                "Updated": "2026-03-01 00:00", "Resolved": "2026-03-01 00:00",
                "Resolution": "Fixed", "Labels": "", "Components": "Backend",
                "Fix Versions": "v1.0", "Sprint": "Sprint 2",
                "Original Estimate": "8.0h", "Remaining Estimate": "",
                "Time Spent": "8.0h", "Age (Days)": 226,
            },
        ]
        import pandas as pd
        self.df = pd.DataFrame(self.test_data)
        self.tmpdir = tempfile.mkdtemp()
        self.output_path = os.path.join(self.tmpdir, "test_report.xlsx")

    def tearDown(self):
        import gc
        gc.collect()  # Force garbage collection to release file handles
        try:
            if os.path.exists(self.output_path):
                os.remove(self.output_path)
        except PermissionError:
            pass  # Windows may still hold the file lock

    def test_excel_file_created(self):
        result = je.generate_excel(self.df, self.output_path)
        self.assertTrue(os.path.exists(result))
        self.assertGreater(os.path.getsize(result), 0)

    def test_excel_has_issues_sheet(self):
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        self.assertIn("Issues", wb.sheetnames)
        wb.close()

    def test_excel_has_summary_sheet(self):
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        self.assertIn("Summary", wb.sheetnames)
        wb.close()

    def test_excel_issues_row_count(self):
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Issues"]
        data_rows = ws.max_row - 1  # Subtract header row
        self.assertEqual(data_rows, 2)
        wb.close()

    def test_excel_issues_column_count(self):
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Issues"]
        self.assertEqual(ws.max_column, len(self.df.columns))
        wb.close()

    def test_excel_freeze_panes(self):
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Issues"]
        self.assertEqual(ws.freeze_panes, "A2")
        wb.close()

    def test_excel_auto_filter(self):
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Issues"]
        self.assertIsNotNone(ws.auto_filter.ref)
        self.assertGreater(len(ws.auto_filter.ref), 0)
        wb.close()

    def test_excel_header_styling(self):
        """Headers should have bold font and fill color."""
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Issues"]
        header_cell = ws.cell(row=1, column=1)
        self.assertTrue(header_cell.font.bold)
        self.assertIsNotNone(header_cell.fill.start_color)
        wb.close()

    def test_summary_contains_total(self):
        """Summary sheet should contain the total issue count."""
        je.generate_excel(self.df, self.output_path)
        import openpyxl
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Summary"]
        # Search for the total count value in the sheet
        found = False
        for row in ws.iter_rows(values_only=True):
            if 2 in row:  # We have 2 test issues
                found = True
                break
        self.assertTrue(found, "Summary sheet should contain total issue count")
        wb.close()


# ---------------------------------------------------------------------------
# 5. API Error Handling Tests (Mocked)
# ---------------------------------------------------------------------------
class TestAPIErrorHandling(unittest.TestCase):
    """Tests for API error handling using mocked HTTP responses."""

    @patch("requests.get")
    def test_fetch_issues_success(self, mock_get):
        """Successful API response returns _IssueProxy list."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {
            "total": 1,
            "issues": [
                {
                    "key": "TEST-1",
                    "fields": {"summary": "Mock issue", "status": {"name": "Open"}},
                }
            ],
        }
        mock_resp.raise_for_status = MagicMock()
        mock_get.return_value = mock_resp

        issues = je.fetch_issues(
            "https://mock.atlassian.net", "test@test.com", "fake-token",
            "project = TEST", limit=1,
        )
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0].key, "TEST-1")

    @patch("requests.get")
    def test_fetch_issues_pagination(self, mock_get):
        """Pagination correctly accumulates multiple batches."""
        page1 = MagicMock()
        page1.status_code = 200
        page1.json.return_value = {
            "total": 3,
            "issues": [
                {"key": f"T-{i}", "fields": {"summary": f"Issue {i}"}}
                for i in range(1, 3)
            ],
        }
        page1.raise_for_status = MagicMock()

        page2 = MagicMock()
        page2.status_code = 200
        page2.json.return_value = {
            "total": 3,
            "issues": [
                {"key": "T-3", "fields": {"summary": "Issue 3"}}
            ],
        }
        page2.raise_for_status = MagicMock()

        mock_get.side_effect = [page1, page2]

        issues = je.fetch_issues(
            "https://mock.atlassian.net", "test@test.com", "fake-token",
            "project = TEST", max_per_page=2,
        )
        self.assertEqual(len(issues), 3)

    @patch("requests.get")
    def test_fetch_issues_limit(self, mock_get):
        """--limit correctly caps the number of results."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {
            "total": 100,
            "issues": [
                {"key": f"T-{i}", "fields": {"summary": f"Issue {i}"}}
                for i in range(5)
            ],
        }
        mock_resp.raise_for_status = MagicMock()
        mock_get.return_value = mock_resp

        issues = je.fetch_issues(
            "https://mock.atlassian.net", "test@test.com", "fake-token",
            "project = TEST", limit=5,
        )
        self.assertLessEqual(len(issues), 5)

    @patch("requests.get")
    def test_fetch_issues_http_401_exits(self, mock_get):
        """HTTP 401 should cause sys.exit."""
        import requests
        mock_resp = MagicMock()
        mock_resp.status_code = 401
        mock_resp.text = "Unauthorized"
        http_error = requests.exceptions.HTTPError(response=mock_resp)
        mock_resp.raise_for_status.side_effect = http_error
        mock_get.return_value = mock_resp

        with self.assertRaises(SystemExit):
            je.fetch_issues(
                "https://mock.atlassian.net", "test@test.com", "bad-token",
                "project = TEST",
            )

    @patch("requests.get")
    def test_fetch_issues_http_404_exits(self, mock_get):
        """HTTP 404 should cause sys.exit."""
        import requests
        mock_resp = MagicMock()
        mock_resp.status_code = 404
        mock_resp.text = "Not Found"
        http_error = requests.exceptions.HTTPError(response=mock_resp)
        mock_resp.raise_for_status.side_effect = http_error
        mock_get.return_value = mock_resp

        with self.assertRaises(SystemExit):
            je.fetch_issues(
                "https://mock.atlassian.net", "test@test.com", "token",
                "project = NONEXIST",
            )

    @patch("requests.get")
    def test_fetch_issues_empty_result(self, mock_get):
        """Empty query result returns empty list."""
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {"total": 0, "issues": []}
        mock_resp.raise_for_status = MagicMock()
        mock_get.return_value = mock_resp

        issues = je.fetch_issues(
            "https://mock.atlassian.net", "test@test.com", "token",
            "project = EMPTY",
        )
        self.assertEqual(len(issues), 0)


# ---------------------------------------------------------------------------
# 6. CLI Argument Parsing Tests
# ---------------------------------------------------------------------------
class TestCLIParsing(unittest.TestCase):
    """Tests for command-line argument parsing."""

    def test_default_args(self):
        with patch("sys.argv", ["jira_extraction.py"]):
            args = je.parse_args()
            self.assertIsNone(args.jql)
            self.assertIsNone(args.template)
            self.assertIsNone(args.limit)
            self.assertIsNone(args.output)
            self.assertEqual(args.config, "config.json")

    def test_jql_arg(self):
        with patch("sys.argv", ["jira_extraction.py", "--jql", "project = DEV"]):
            args = je.parse_args()
            self.assertEqual(args.jql, "project = DEV")

    def test_template_arg(self):
        with patch("sys.argv", ["jira_extraction.py", "--template", "open_bugs"]):
            args = je.parse_args()
            self.assertEqual(args.template, "open_bugs")

    def test_limit_arg(self):
        with patch("sys.argv", ["jira_extraction.py", "--limit", "50"]):
            args = je.parse_args()
            self.assertEqual(args.limit, 50)

    def test_output_arg(self):
        with patch("sys.argv", ["jira_extraction.py", "--output", "my_report.xlsx"]):
            args = je.parse_args()
            self.assertEqual(args.output, "my_report.xlsx")


# ---------------------------------------------------------------------------
# 7. Live Integration Smoke Test
# ---------------------------------------------------------------------------
class TestLiveIntegration(unittest.TestCase):
    """Live smoke test – only runs when real credentials are available."""

    def test_live_jira_connection(self):
        """Verify we can connect to Jira and fetch at least 1 issue."""
        # Load real credentials
        from dotenv import load_dotenv
        env_path = Path(__file__).resolve().parent.parent / ".env"
        load_dotenv(dotenv_path=str(env_path))

        url = os.getenv("JIRA_URL")
        email = os.getenv("JIRA_EMAIL")
        token = os.getenv("JIRA_API_TOKEN")

        if not all([url, email, token]):
            self.skipTest("Live credentials not available")

        # Attempt to fetch 1 issue
        issues = je.fetch_issues(url, email, token,
                                 "project IS NOT EMPTY ORDER BY created DESC",
                                 limit=1)
        self.assertGreaterEqual(len(issues), 1)
        self.assertIsNotNone(issues[0].key)

    def test_live_end_to_end_report(self):
        """Full pipeline: fetch → transform → Excel generation."""
        from dotenv import load_dotenv
        env_path = Path(__file__).resolve().parent.parent / ".env"
        load_dotenv(dotenv_path=str(env_path))

        url = os.getenv("JIRA_URL")
        email = os.getenv("JIRA_EMAIL")
        token = os.getenv("JIRA_API_TOKEN")

        if not all([url, email, token]):
            self.skipTest("Live credentials not available")

        issues = je.fetch_issues(url, email, token,
                                 "project IS NOT EMPTY ORDER BY created DESC",
                                 limit=3)
        if not issues:
            self.skipTest("No issues returned from Jira")

        df = je.issues_to_dataframe(issues)
        self.assertGreater(len(df), 0)

        tmpdir = tempfile.mkdtemp()
        output_path = os.path.join(tmpdir, "live_test_report.xlsx")
        result = je.generate_excel(df, output_path)
        self.assertTrue(os.path.exists(result))
        self.assertGreater(os.path.getsize(result), 0)

        # Cleanup
        os.remove(output_path)


# ---------------------------------------------------------------------------
# 8. File Structure Validation
# ---------------------------------------------------------------------------
class TestProjectStructure(unittest.TestCase):
    """Tests that required project files exist."""

    PROJECT_ROOT = Path(__file__).resolve().parent.parent

    def test_main_script_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "jira_extraction.py").exists())

    def test_config_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "config.json").exists())

    def test_requirements_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "requirements.txt").exists())

    def test_env_example_exists(self):
        self.assertTrue((self.PROJECT_ROOT / ".env.example").exists())

    def test_gitignore_exists(self):
        self.assertTrue((self.PROJECT_ROOT / ".gitignore").exists())

    def test_output_directory_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "output").exists())

    def test_dockerfile_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "Dockerfile").exists())

    def test_docker_compose_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "docker-compose.yml").exists())

    def test_nomad_job_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "nomad" / "jira-extraction.nomad").exists())

    def test_deploy_ps1_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "deploy.ps1").exists())

    def test_deploy_sh_exists(self):
        self.assertTrue((self.PROJECT_ROOT / "deploy.sh").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
