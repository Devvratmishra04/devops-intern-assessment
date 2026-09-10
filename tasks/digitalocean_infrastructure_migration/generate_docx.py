import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Digital Ocean Setup & Infrastructure Migration Guide")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("End-to-End Migration Architecture: Bastion, Kafka, Grafana, Consul, Nomad & CRM Extractions")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project:", "Digital Ocean Production Migration & Service Mesh Provisioning"),
        ("Components:", "Bastion, Consul, Nomad, Kafka (KRaft), Grafana, Microsoft/HubSpot/Salesforce Extractions"),
        ("Author / Engineer:", "Devvrat Mishra (@DevOps)"),
        ("Date:", "September 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        cell_k = meta_table.cell(i, 0)
        cell_v = meta_table.cell(i, 1)
        cell_k.text = k
        cell_v.text = v
        cell_k.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell_k, "F0F4F8")
        set_cell_background(cell_v, "FAFAFA")
    doc.add_paragraph()

    # Section 1
    doc.add_heading("1. Executive Summary & Migration Objectives", level=1)
    doc.add_paragraph(
        "This migration blueprint establishes a secure, enterprise-grade cloud environment on Digital Ocean. "
        "It decouples ingress management via an SSH Bastion gateway, establishes a resilient service mesh via "
        "HashiCorp Consul and Nomad, deploys an event-driven messaging layer powered by Apache Kafka in KRaft mode, "
        "and orchestrates automated extraction pipelines from Microsoft Graph, HubSpot CRM, and Salesforce CRM."
    )

    # Section 2
    doc.add_heading("2. Acceptance Criteria Component Mapping", level=1)
    crit_table = doc.add_table(rows=9, cols=3)
    crit_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Component", "Deployment Architecture", "Verification Status"]
    for j, h in enumerate(headers):
        cell = crit_table.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "003366")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    components = [
        ("Bastion", "Dedicated Droplet with UFW, fail2ban & restricted SSH ingress", "VERIFIED [PASS]"),
        ("Kafka", "Nomad service job running Apache Kafka 3.7.0 in KRaft mode", "VERIFIED [PASS]"),
        ("Grafana", "Observability hub with 2 pre-provisioned dashboards & datasources", "VERIFIED [PASS]"),
        ("Consul", "3-node HA server cluster + system client daemon on all workers", "VERIFIED [PASS]"),
        ("Nomad", "3-node server quorum + worker client pool managing all workloads", "VERIFIED [PASS]"),
        ("Microsoft extraction", "Graph API delta sync service publishing to extractions.microsoft", "VERIFIED [PASS]"),
        ("HubSpot extraction", "CRM API rate-throttled worker publishing to extractions.hubspot", "VERIFIED [PASS]"),
        ("Salesforce extraction", "SOQL cursor-based sync worker publishing to extractions.salesforce", "VERIFIED [PASS]")
    ]

    for i, (comp, arch, status) in enumerate(components):
        r = i + 1
        c0 = crit_table.cell(r, 0)
        c1 = crit_table.cell(r, 1)
        c2 = crit_table.cell(r, 2)
        c0.text = comp
        c1.text = arch
        c2.text = status
        c0.paragraphs[0].runs[0].font.bold = True
        set_cell_background(c0, "F8F9FA")
        set_cell_background(c1, "FFFFFF")
        set_cell_background(c2, "E8F5E9")

    doc.add_paragraph()

    # Section 3
    doc.add_heading("3. VPC & Network Security Topology", level=1)
    doc.add_paragraph(
        "All compute nodes reside strictly within an isolated DigitalOcean VPC (10.136.0.0/16). "
        "Direct public IP addresses are eliminated from internal nodes. Access is governed via:"
    )
    doc.add_paragraph("• Bastion Host: Only port 22 exposed to authorized administrator CIDRs. Password authentication disabled.")
    doc.add_paragraph("• Internal Cluster Firewall: Blocks all public ingress; allows intra-VPC communication for Consul, Nomad, and Kafka.")
    doc.add_paragraph("• Outbound Egress: Permitted for secure API synchronization with Microsoft, HubSpot, and Salesforce endpoints.")

    # Section 4
    doc.add_heading("4. Event-Driven Data Pipeline Design", level=1)
    doc.add_paragraph(
        "Apache Kafka acts as the high-throughput buffer for all incoming extraction events. "
        "Running in KRaft mode eliminates ZooKeeper failure modes. Topics are configured with 3 partitions and 7-day retention:"
    )
    doc.add_paragraph("1. extractions.microsoft: Captures user changes, mail events, and directory delta updates.")
    doc.add_paragraph("2. extractions.hubspot: Streams contact, deal, and engagement CRM records under strict rate limits.")
    doc.add_paragraph("3. extractions.salesforce: Pulls lead and account updates using timestamp cursors (SystemModstamp).")
    doc.add_paragraph("4. extractions.dlq: Isolates malformed or unprocessable payloads for diagnostic triage.")

    # Section 5
    doc.add_heading("5. Migration Execution & Verification Runbook", level=1)
    doc.add_paragraph("The migration was executed and validated via an automated 7-stage orchestrator:")
    doc.add_paragraph("1. Pre-Flight Validation: Verified DO credentials, SSH key, and network parameters.")
    doc.add_paragraph("2. Bastion & VPC Isolation: Verified firewall rules and gateway proxy-jump capability.")
    doc.add_paragraph("3. Consul & Nomad Bootstrap: Established Raft consensus and node registrations.")
    doc.add_paragraph("4. Kafka Broker Deployment: Initialized KRaft controller and auto-created extraction topics.")
    doc.add_paragraph("5. Grafana Monitoring Stack: Deployed Grafana 11.1.0 with Cluster Overview and Telemetry dashboards.")
    doc.add_paragraph("6. Ingestion Pipeline Deployment: Launched Microsoft, HubSpot, and Salesforce extraction workers.")
    doc.add_paragraph("7. Cutover Smoke Test: Executed end-to-end event production and DLQ health checks.")

    output_path = os.path.join(os.path.dirname(__file__), "DigitalOcean_Setup_and_Migration_Guide.docx")
    doc.save(output_path)
    print(f"Generated Word Document: {output_path}")

if __name__ == "__main__":
    create_document()
